from __future__ import annotations

import asyncio
import base64
import io
import json
import logging
import os
import time
from pathlib import Path
from typing import Any, Dict, Optional, Callable

import openai
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field

from scoutai.backend.config import validate_config, list_models, NEBIUS_API_KEY, NEBIUS_BASE_URL, TAVILY_API_KEY
from scoutai.backend.pipeline import PipelineRun, PipelineProgress
from scoutai.backend.searcher import extract_urls
from scoutai.backend.models import PlanRequest, NewQuestionRequest, PipelineStatus, ReportData

logger = logging.getLogger(__name__)

app = FastAPI(title="ScoutAI", description="Multi-model deep research agent", version="0.1.0")

# --- Static / Templates ---
PROJECT_DIR = Path(__file__).parent.parent
FRONTEND_DIR = PROJECT_DIR / "frontend"
STATIC_DIR = FRONTEND_DIR / "static"
TEMPLATE_DIR = FRONTEND_DIR / "templates"

if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

templates = Jinja2Templates(directory=str(TEMPLATE_DIR) if TEMPLATE_DIR.exists() else str(PROJECT_DIR / "templates"))

# --- Global state ---
active_runs: Dict[str, PipelineRun] = {}


@app.on_event("startup")
async def startup_check():
    logger.info("ScoutAI startup check")
    if validate_config():
        logger.info(f"Configured models: {list_models()}")
    else:
        logger.warning("Startup config validation failed — running in demo mode")


@app.get("/health", response_model=Dict[str, str])
async def health():
    return {"status": "ok", "message": "ScoutAI is running"}


@app.get("/api/status")
async def api_status():
    missing = []
    if not NEBIUS_API_KEY:
        missing.append("NEBIUS_API_KEY")
    if not TAVILY_API_KEY:
        missing.append("TAVILY_API_KEY")
    return {
        "demo": bool(missing),
        "missing": missing,
        "models": list_models(),
    }


@app.get("/api/contract")
async def api_contract():
    return {
        "events": [
            {"name": "plan_start", "description": "Planning step started", "fields": ["step", "model"]},
            {"name": "plan_done", "description": "Planning step completed", "fields": ["step", "model", "sub_questions_count"]},
            {"name": "plan_error", "description": "Planning step failed", "fields": ["step", "model", "error"]},
            {"name": "search_start", "description": "Search step started", "fields": ["step", "model", "query_count"]},
            {"name": "search_done", "description": "Search step completed", "fields": ["step", "model", "sources_count"]},
            {"name": "search_error", "description": "Search step failed", "fields": ["step", "model", "error"]},
            {"name": "extract_start", "description": "Extraction step started", "fields": ["step", "model", "item_count"]},
            {"name": "extract_done", "description": "Extraction step completed", "fields": ["step", "model", "bullets_count"]},
            {"name": "extract_error", "description": "Extraction step failed", "fields": ["step", "model", "error"]},
            {"name": "verify_start", "description": "Verification step started", "fields": ["step", "model", "claim_count"]},
            {"name": "verify_done", "description": "Verification step completed", "fields": ["step", "model", "claims_verified"]},
            {"name": "verify_error", "description": "Verification step failed", "fields": ["step", "model", "error"]},
            {"name": "write_start", "description": "Report writing step started", "fields": ["step", "model"]},
            {"name": "write_done", "description": "Report writing step completed", "fields": ["step", "model", "report_id"]},
            {"name": "write_error", "description": "Writing step failed", "fields": ["step", "model", "error"]},
            {"name": "complete", "description": "Pipeline complete with final report", "fields": ["report", "citations", "total_time"]},
        ]
    }


@app.get("/sse/{run_id}")
async def sse_stream(run_id: str):
    run = active_runs.get(run_id)
    if not run:
        return JSONResponse({"type": "error", "data": {"message": "Run not found"}}, status_code=404)

    async def event_stream():
        def complete_event():
            result = run.write_result or {}
            return {
                "type": "complete",
                "data": {
                    "report": result.get("report", ""),
                    "citations": result.get("citations", []),
                    "total_time": round(time.time() - run.start_time, 2),
                },
            }

        if run.write_result:
            yield f"data: {json.dumps(complete_event())}\n\n"
            return

        last_key = None
        # Stream real pipeline progress; allow up to ~5 minutes for thorough runs
        for _ in range(3000):
            if run.write_result:
                yield f"data: {json.dumps(complete_event())}\n\n"
                return

            prog = getattr(run, "progress", None)
            if prog is not None:
                key = (prog.step, prog.done, prog.step_description)
                if key != last_key:
                    last_key = key
                    yield f"data: {json.dumps({'type': 'progress', 'data': {'step': prog.step, 'status': 'done' if prog.done else 'running', 'model': prog.model, 'step_description': prog.step_description}})}\n\n"
            elif last_key is None:
                last_key = ("plan", False, "Planning research questions...")
                yield f"data: {json.dumps({'type': 'progress', 'data': {'step': 'plan', 'status': 'running', 'model': run.model_fast, 'step_description': 'Planning research questions...'}})}\n\n"

            await asyncio.sleep(0.1)

        yield f"data: {json.dumps({'type': 'error', 'data': {'message': 'Pipeline timeout'}})}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@app.post("/api/run")
async def api_run(request: NewQuestionRequest, background_tasks: Any = None):
    run_id = f"run_{int(time.time() * 1000)}"
    if run_id in active_runs:
        del active_runs[run_id]
    run = PipelineRun(
        question=request.question,
        on_progress=lambda prog: None,
        on_complete=lambda r: None,
        on_error=lambda err: logger.error(f"Pipeline error: {err}"),
        seed_urls=extract_urls(request.question),
    )
    model_choice = (request.model or "").strip()
    if model_choice:
        configured = list_models()
        run.model_strong = {
            "nemotron-nano": configured.get("fast"),
            "nemotron-super": configured.get("mid"),
            "nemotron-ultra": configured.get("strong"),
        }.get(model_choice) or run.model_strong
    active_runs[run_id] = run
    import asyncio
    async def _run_and_store():
        try:
            result = await run.run()
            active_runs[run_id].write_result = result
        except Exception as e:
            logger.exception("Background pipeline error")
            active_runs[run_id].write_result = {"error": str(e)}
    asyncio.create_task(_run_and_store())
    return {"run_id": run_id, "status": "started"}


@app.get("/api/run/{run_id}")
async def api_run_status(run_id: str):
    run = active_runs.get(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    total_time = round(time.time() - run.start_time, 2) if hasattr(run, 'start_time') else 0
    return {"run_id": run_id, "status": "completed" if run.write_result else "running", "total_time": total_time, "has_report": run.write_result is not None}


@app.post("/api/demo")
async def api_demo(request: NewQuestionRequest):
    question = (request.question or "").strip()
    if len(question) > 220:
        question = question[:217].rstrip() + "..."
    topic = question or "your research question"
    report = (
        f"EXECUTIVE SUMMARY:\n"
        f"This is a demo report for \u201c{topic}\u201d. It is generated from cached data so you can explore "
        f"the interface; add your API keys to .env to run the live 5-step research pipeline.\n\n"
        "KEY FINDINGS:\n"
        "1. ScoutAI plans sub-questions, searches the web with Tavily, and extracts evidence from the top sources.\n"
        "2. Each extracted claim is cross-checked for confidence before the writer assembles the final brief with citations.\n\n"
        "CONFIDENT OR UNCERTAIN INFO:\n"
        "- Confirmed: Demo mode is active because the Nebius and/or Tavily API keys are missing.\n"
        "- Uncertain/Conflicting: The findings shown here are placeholder content, not live research.\n\n"
        "CITATIONS:\n"
        "[1] URL: https://example.com\n"
    )
    return {"report": report}




IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


async def _describe_image(data: bytes, content_type: str, name: str) -> str:
    mime = content_type or "image/jpeg"
    if not mime.startswith("image/"):
        mime = "image/jpeg"
    encoded = base64.b64encode(data).decode("ascii")

    def call_model() -> str:
        client = openai.OpenAI(
            base_url=NEBIUS_BASE_URL,
            api_key=NEBIUS_API_KEY,
            timeout=90.0,
            max_retries=0,
        )
        response = client.chat.completions.create(
            model="nvidia/Nemotron-3.5-Super-VL",
            messages=[{
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "You are reading an uploaded image for a research assistant. "
                            "Extract all visible text verbatim, then briefly describe the key "
                            "content, chart, table, or diagram. Be concise but complete."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{encoded}"},
                    },
                ],
            }],
            temperature=0.1,
            max_tokens=1200,
        )
        message = response.choices[0].message
        return (message.content or "").strip()

    return await asyncio.wait_for(asyncio.to_thread(call_model), timeout=120)


@app.post("/api/extract-file")
async def api_extract_file(file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > 12 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is larger than 12 MB")

    name = file.filename or "file"
    ext = Path(name).suffix.lower()
    text = ""

    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:30])
        elif ext == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif ext in IMAGE_EXTENSIONS:
            text = await _describe_image(data, file.content_type or "", name)
        else:
            text = data.decode("utf-8", errors="ignore")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read {ext or 'file'}: {exc}")

    text = text.strip()
    if not text:
        raise HTTPException(status_code=422, detail="No readable text found in this file")
    return {"name": name, "text": text[:20000]}


@app.get("/", response_class=HTMLResponse)
async def root(request: Request):
    return templates.TemplateResponse("index.html", {"request": request, "title": "ScoutAI"})

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
